import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook

from catalog.models import Level
from crm.models import (
    InformationSource,
    School,
    SchoolEducationalService,
    SchoolImportBatch,
    SchoolImportRow,
    SchoolPopulationRecord,
)


HEADER_ALIASES = {
    "codigo_modular": "modular_code",
    "codigo_de_institucion": "institution_code",
    "codigo_institucion": "institution_code",
    "nombre_de_ie": "school_name",
    "nombre_ie": "school_name",
    "nombre_de_la_ie": "school_name",
    "nivel_modalidad": "level_name",
    "nivel": "level_name",
    "dependencia": "dependency",
    "direccion": "address",
    "departamento": "department",
    "provincia": "province",
    "distrito": "district",
    "alumnos": "students",
}

REQUIRED_COLUMNS = {
    "institution_code": "Código de institución",
    "school_name": "Nombre de IE",
    "level_name": "Nivel/Modalidad",
}


def _normalize_header(value):
    text = str(value or "").strip().lower()
    text = "".join(
        character
        for character in unicodedata.normalize("NFKD", text)
        if not unicodedata.combining(character)
    )
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text


def _clean_text(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()

    return str(value).strip()


def _clean_code(value):
    text = _clean_text(value)

    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]

    return text


def _parse_students(value):
    if value in (None, ""):
        return None

    if isinstance(value, bool):
        raise ValueError("La cantidad de alumnos debe ser numérica.")

    try:
        number = int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError) as error:
        raise ValueError(
            "La cantidad de alumnos debe ser numérica."
        ) from error

    if number < 0:
        raise ValueError(
            "La cantidad de alumnos no puede ser negativa."
        )

    return number


def _resolve_level(raw_value):
    raw_text = _clean_text(raw_value)
    normalized = _normalize_header(raw_text)

    candidates = []

    if "inicial" in normalized:
        candidates.append("inicial")
    if "primaria" in normalized:
        candidates.append("primaria")
    if "secundaria" in normalized:
        candidates.append("secundaria")

    candidates.append(slugify(raw_text))

    for slug in dict.fromkeys(candidates):
        if not slug:
            continue

        level = Level.objects.filter(
            slug__iexact=slug,
            is_active=True,
        ).first()

        if level is not None:
            return level

    lowered = raw_text.lower()
    return (
        Level.objects
        .filter(name__iexact=lowered, is_active=True)
        .first()
    )


def _sheet_and_columns(file_path):
    workbook = load_workbook(file_path, data_only=True)

    if "Instituciones" in workbook.sheetnames:
        sheet = workbook["Instituciones"]
    else:
        sheet = workbook.active

    columns = {}

    for index, cell in enumerate(sheet[1]):
        normalized = _normalize_header(cell.value)
        canonical = HEADER_ALIASES.get(normalized)

        if canonical and canonical not in columns:
            columns[canonical] = index

    missing = [
        display_name
        for canonical, display_name in REQUIRED_COLUMNS.items()
        if canonical not in columns
    ]

    if missing:
        raise ValueError(
            "Faltan columnas obligatorias: "
            + ", ".join(missing)
            + "."
        )

    return sheet, columns


def _cell(row, columns, key):
    index = columns.get(key)

    if index is None or index >= len(row):
        return None

    return row[index]


def _read_rows(file_path):
    sheet, columns = _sheet_and_columns(file_path)
    rows = []

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        if all(value in (None, "") for value in row):
            continue

        data = {
            "institution_code": _clean_code(
                _cell(row, columns, "institution_code")
            ),
            "modular_code": _clean_code(
                _cell(row, columns, "modular_code")
            ),
            "school_name": _clean_text(
                _cell(row, columns, "school_name")
            ),
            "level_name": _clean_text(
                _cell(row, columns, "level_name")
            ),
            "dependency": _clean_text(
                _cell(row, columns, "dependency")
            ),
            "address": _clean_text(
                _cell(row, columns, "address")
            ),
            "department": _clean_text(
                _cell(row, columns, "department")
            ),
            "province": _clean_text(
                _cell(row, columns, "province")
            ),
            "district": _clean_text(
                _cell(row, columns, "district")
            ),
            "students_raw": _cell(row, columns, "students"),
        }

        rows.append((row_number, data))

    return rows


def _base_signature(data):
    return (
        data.get("school_name", "").strip().casefold(),
        data.get("department", "").strip().casefold(),
        data.get("province", "").strip().casefold(),
        data.get("district", "").strip().casefold(),
    )


def _preview_error_batch(batch, message):
    SchoolImportRow.objects.create(
        batch=batch,
        row_number=1,
        action=SchoolImportRow.Action.ERROR,
        errors=[message],
        data={},
    )
    batch.total_errors = 1
    batch.status = SchoolImportBatch.Status.ERROR
    batch.save(
        update_fields=[
            "total_errors",
            "status",
            "updated_at",
        ]
    )
    return batch


def create_school_import_preview(
    *,
    file,
    population_year,
    actor,
):
    batch = SchoolImportBatch.objects.create(
        file=file,
        population_year=population_year,
        created_by=actor,
        status=SchoolImportBatch.Status.PENDING,
    )

    try:
        batch.file.open("rb")
        excel_rows = _read_rows(batch.file)
    except Exception as error:
        return _preview_error_batch(batch, str(error))
    finally:
        batch.file.close()

    seen_school_level = set()
    seen_modular_codes = {}
    school_signatures = {}
    school_actions = {}
    valid_school_codes = set()
    error_count = 0

    for row_number, data in excel_rows:
        errors = []

        institution_code = data["institution_code"]
        modular_code = data["modular_code"]
        school_name = data["school_name"]
        level_raw = data["level_name"]

        if not institution_code:
            errors.append(
                "El Código de institución es obligatorio."
            )

        if not school_name:
            errors.append(
                "El Nombre de IE es obligatorio."
            )

        if not level_raw:
            errors.append(
                "El Nivel/Modalidad es obligatorio."
            )

        level = _resolve_level(level_raw) if level_raw else None

        if level_raw and level is None:
            errors.append(
                (
                    f"No se reconoce el nivel '{level_raw}'. "
                    "Usa Inicial, Primaria o Secundaria."
                )
            )

        try:
            students = _parse_students(data["students_raw"])
        except ValueError as error:
            students = None
            errors.append(str(error))

        data["students"] = students
        data.pop("students_raw", None)

        if level is not None:
            data["level_id"] = level.id
            data["level_display"] = level.name

        if institution_code and level is not None:
            school_level_key = (
                institution_code.casefold(),
                level.id,
            )

            if school_level_key in seen_school_level:
                errors.append(
                    (
                        "El mismo colegio y nivel aparecen más de una "
                        "vez en el Excel."
                    )
                )

            seen_school_level.add(school_level_key)

        if modular_code:
            modular_key = modular_code.casefold()
            previous_row = seen_modular_codes.get(modular_key)

            if previous_row is not None:
                errors.append(
                    (
                        f"El Código modular se repite en la fila "
                        f"{previous_row}."
                    )
                )
            else:
                seen_modular_codes[modular_key] = row_number

            conflicting_service = (
                SchoolEducationalService.objects
                .filter(modular_code__iexact=modular_code)
                .exclude(
                    school__institution_code__iexact=(
                        institution_code
                    )
                )
                .first()
            )

            if conflicting_service is not None:
                errors.append(
                    (
                        "El Código modular ya pertenece a otra "
                        "institución registrada."
                    )
                )

        if institution_code:
            signature = _base_signature(data)
            previous_signature = school_signatures.get(
                institution_code.casefold()
            )

            if (
                previous_signature is not None
                and previous_signature != signature
            ):
                errors.append(
                    (
                        "Los datos principales de esta institución no "
                        "coinciden entre sus filas."
                    )
                )
            else:
                school_signatures[
                    institution_code.casefold()
                ] = signature

        existing_school = None
        if institution_code:
            existing_school = School.objects.filter(
                institution_code__iexact=institution_code
            ).first()

        action = (
            SchoolImportRow.Action.UPDATE
            if existing_school is not None
            else SchoolImportRow.Action.NEW
        )

        if errors:
            action = SchoolImportRow.Action.ERROR
            error_count += 1
        elif institution_code:
            valid_school_codes.add(institution_code.casefold())
            school_actions.setdefault(
                institution_code.casefold(),
                (
                    SchoolImportRow.Action.UPDATE
                    if existing_school is not None
                    else SchoolImportRow.Action.NEW
                ),
            )

        SchoolImportRow.objects.create(
            batch=batch,
            row_number=row_number,
            institution_code=institution_code,
            modular_code=modular_code,
            school_name=school_name,
            level_name=level_raw,
            action=action,
            errors=errors,
            data=data,
        )

    if not excel_rows:
        return _preview_error_batch(
            batch,
            "El archivo no contiene instituciones para importar.",
        )

    total_new = sum(
        1
        for action in school_actions.values()
        if action == SchoolImportRow.Action.NEW
    )
    total_updated = sum(
        1
        for action in school_actions.values()
        if action == SchoolImportRow.Action.UPDATE
    )

    batch.total_rows = len(excel_rows)
    batch.total_schools = len(valid_school_codes)
    batch.total_new = total_new
    batch.total_updated = total_updated
    batch.total_errors = error_count
    batch.status = (
        SchoolImportBatch.Status.VALIDATED
        if error_count == 0
        else SchoolImportBatch.Status.ERROR
    )
    batch.save(
        update_fields=[
            "total_rows",
            "total_schools",
            "total_new",
            "total_updated",
            "total_errors",
            "status",
            "updated_at",
        ]
    )

    return batch


def _update_school_from_rows(
    *,
    school,
    rows,
    actor,
):
    first = rows[0].data

    school.name = first["school_name"]

    for field in (
        "dependency",
        "address",
        "department",
        "province",
        "district",
    ):
        value = first.get(field, "")
        if value:
            setattr(school, field, value)

    school.is_active = True

    modular_codes = {
        row.data.get("modular_code")
        for row in rows
        if row.data.get("modular_code")
    }
    school.modular_code = (
        next(iter(modular_codes))
        if len(modular_codes) == 1
        else ""
    )

    student_values = [
        row.data.get("students")
        for row in rows
        if row.data.get("students") is not None
    ]
    if student_values:
        school.estimated_students = sum(student_values)

    if school.created_by_id is None:
        school.created_by = actor

    school.full_clean()
    school.save()


def _upsert_population(
    *,
    service,
    student_count,
    population_year,
    actor,
    source_detail,
):
    if student_count is None:
        return

    current = (
        service.population_records
        .filter(is_current=True)
        .order_by("-year", "-created_at")
        .first()
    )

    if current is not None and current.year == population_year:
        current.student_count = student_count
        current.source = InformationSource.IMPORT
        current.source_detail = source_detail
        current.recorded_by = actor
        current.save(
            update_fields=[
                "student_count",
                "source",
                "source_detail",
                "recorded_by",
                "updated_at",
            ]
        )
        return

    service.population_records.filter(
        is_current=True
    ).update(is_current=False)

    SchoolPopulationRecord.objects.create(
        service=service,
        year=population_year,
        student_count=student_count,
        source=InformationSource.IMPORT,
        source_detail=source_detail,
        is_current=True,
        recorded_by=actor,
    )


@transaction.atomic
def confirm_school_import(*, batch_id, actor):
    batch = (
        SchoolImportBatch.objects
        .select_for_update()
        .prefetch_related("rows")
        .get(pk=batch_id)
    )

    if batch.status == SchoolImportBatch.Status.IMPORTED:
        raise ValueError(
            "Esta importación ya fue confirmada."
        )

    if batch.total_errors > 0:
        raise ValueError(
            "Corrige los errores del archivo antes de confirmar."
        )

    if batch.status != SchoolImportBatch.Status.VALIDATED:
        raise ValueError(
            "La importación todavía no está lista para confirmar."
        )

    grouped_rows = defaultdict(list)

    for row in batch.rows.all():
        grouped_rows[row.institution_code.casefold()].append(row)

    source_detail = (
        f"Excel {Path(batch.file.name).name}"
    )

    for rows in grouped_rows.values():
        rows.sort(key=lambda item: item.row_number)
        first = rows[0]
        institution_code = first.institution_code

        school = School.objects.filter(
            institution_code__iexact=institution_code
        ).first()

        if school is None:
            school = School(
                institution_code=institution_code,
                name=first.school_name,
                created_by=actor,
            )

        _update_school_from_rows(
            school=school,
            rows=rows,
            actor=actor,
        )

        imported_levels = []

        for row in rows:
            data = row.data
            level = Level.objects.get(pk=data["level_id"])
            imported_levels.append(level)

            service, created = (
                SchoolEducationalService.objects.get_or_create(
                    school=school,
                    level=level,
                    defaults={
                        "modular_code": (
                            data.get("modular_code") or None
                        ),
                        "modality": data.get("level_name", ""),
                        "is_active": True,
                        "created_by": actor,
                    },
                )
            )

            if not created:
                service.modular_code = (
                    data.get("modular_code") or None
                )
                service.modality = data.get("level_name", "")
                service.is_active = True
                service.full_clean()
                service.save(
                    update_fields=[
                        "modular_code",
                        "modality",
                        "is_active",
                        "updated_at",
                    ]
                )

            _upsert_population(
                service=service,
                student_count=data.get("students"),
                population_year=batch.population_year,
                actor=actor,
                source_detail=source_detail,
            )

            row.processed = True
            row.save(
                update_fields=[
                    "processed",
                    "updated_at",
                ]
            )

        school.levels.add(*imported_levels)

    batch.status = SchoolImportBatch.Status.IMPORTED
    batch.save(
        update_fields=[
            "status",
            "updated_at",
        ]
    )

    return batch
