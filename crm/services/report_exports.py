from io import BytesIO

from django.contrib.auth import get_user_model
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from accounts.permissions import usuario_es_administrador
from crm.models import Campaign, CommercialTeam
from crm.selectors import supervised_team_ids

from .reports import (
    build_activity_commercial_report,
    build_advisor_commercial_report,
    build_editorial_commercial_report,
    build_opportunity_commercial_report,
    build_school_commercial_report,
)


BLACK = "111827"
RED = "B91C1C"
LIGHT_RED = "FEE2E2"
LIGHT_GRAY = "F3F4F6"
WHITE = "FFFFFF"
AMBER = "FEF3C7"

THIN_GRAY = Side(style="thin", color="D1D5DB")
TABLE_BORDER = Border(
    left=THIN_GRAY,
    right=THIN_GRAY,
    top=THIN_GRAY,
    bottom=THIN_GRAY,
)


def _safe_local_datetime(value):
    if value is None:
        return None

    if timezone.is_aware(value):
        value = timezone.localtime(value)

    return value.replace(tzinfo=None)


def _advisor_name(user):
    if user is None:
        return ""

    return user.get_full_name().strip() or user.get_username()


def _filter_labels(
    *,
    user,
    campaign_id=None,
    team_id=None,
    owner_id=None,
):
    campaign_label = "Todas las campañas"
    team_label = "Todos los equipos"
    owner_label = "Todos los asesores"

    if campaign_id is not None:
        campaign_label = (
            Campaign.objects
            .filter(pk=campaign_id)
            .values_list("name", flat=True)
            .first()
            or "Campaña no disponible"
        )

    if team_id is not None:
        teams = CommercialTeam.objects.filter(
            pk=team_id,
            is_active=True,
        )

        if not usuario_es_administrador(user):
            teams = teams.filter(
                pk__in=supervised_team_ids(user)
            )

        team_label = (
            teams.values_list("name", flat=True).first()
            or "Equipo fuera de alcance"
        )

    if owner_id is not None:
        advisor_report = build_advisor_commercial_report(
            user=user,
            campaign_id=campaign_id,
            team_id=team_id,
            owner_id=owner_id,
            date_from=timezone.localdate(),
            date_to=timezone.localdate(),
        )
        advisor_row = next(
            (
                row
                for row in advisor_report["advisors"]
                if row["advisor_id"] == owner_id
            ),
            None,
        )

        if advisor_row is not None:
            owner_label = advisor_row["advisor_name"]
        else:
            User = get_user_model()
            owner = User.objects.filter(pk=owner_id).first()
            owner_label = (
                "Asesor fuera de alcance"
                if owner is not None
                else "Asesor no disponible"
            )

    return {
        "campaign": campaign_label,
        "team": team_label,
        "owner": owner_label,
    }


def _style_title(ws, title, subtitle=None):
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = title
    title_cell.fill = PatternFill("solid", fgColor=BLACK)
    title_cell.font = Font(
        color=WHITE,
        bold=True,
        size=16,
    )
    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )
    ws.row_dimensions[1].height = 28

    if subtitle:
        ws.merge_cells("A2:H2")
        subtitle_cell = ws["A2"]
        subtitle_cell.value = subtitle
        subtitle_cell.font = Font(
            color="4B5563",
            italic=True,
            size=10,
        )
        subtitle_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )


def _write_metadata(
    ws,
    *,
    start_row,
    labels,
    date_from,
    date_to,
):
    metadata = [
        ("Campaña", labels["campaign"]),
        ("Equipo", labels["team"]),
        ("Asesor", labels["owner"]),
        (
            "Periodo",
            f"{date_from.strftime('%d/%m/%Y')} - "
            f"{date_to.strftime('%d/%m/%Y')}",
        ),
        (
            "Generado",
            timezone.localtime().strftime("%d/%m/%Y %H:%M"),
        ),
    ]

    row = start_row

    for label, value in metadata:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).font = Font(
            bold=True,
            color=BLACK,
        )
        ws.cell(row=row, column=2, value=value)
        row += 1

    return row


def _write_table(
    ws,
    *,
    start_row,
    headers,
    rows,
    widths=None,
    currency_columns=None,
    percent_columns=None,
    date_columns=None,
):
    currency_columns = set(currency_columns or [])
    percent_columns = set(percent_columns or [])
    date_columns = set(date_columns or [])

    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(
            row=start_row,
            column=column_index,
            value=header,
        )
        cell.fill = PatternFill("solid", fgColor=BLACK)
        cell.font = Font(
            color=WHITE,
            bold=True,
            size=10,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = TABLE_BORDER

    for row_offset, values in enumerate(rows, start=1):
        excel_row = start_row + row_offset

        for column_index, value in enumerate(values, start=1):
            cell = ws.cell(
                row=excel_row,
                column=column_index,
                value=value,
            )
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            if column_index in currency_columns:
                cell.number_format = '"S/" #,##0.00'
            elif column_index in percent_columns:
                cell.number_format = '0.0"%"'
            elif column_index in date_columns:
                cell.number_format = "dd/mm/yyyy hh:mm"

    if widths:
        for column_index, width in enumerate(widths, start=1):
            ws.column_dimensions[
                get_column_letter(column_index)
            ].width = width

    ws.auto_filter.ref = (
        f"A{start_row}:"
        f"{get_column_letter(len(headers))}"
        f"{start_row + len(rows)}"
    )
    ws.freeze_panes = f"A{start_row + 1}"

    return start_row + len(rows) + 1


def _add_summary_sheet(
    wb,
    *,
    advisor_report,
    labels,
    date_from,
    date_to,
):
    ws = wb.active
    ws.title = "01_Resumen"

    _style_title(
        ws,
        "BOOK EXPRESS - REPORTE COMERCIAL CRM",
        "Resumen ejecutivo del alcance seleccionado",
    )
    next_row = _write_metadata(
        ws,
        start_row=4,
        labels=labels,
        date_from=date_from,
        date_to=date_to,
    )

    summary = advisor_report["summary"]
    unassigned = advisor_report["unassigned"]

    next_row += 1
    ws.cell(
        row=next_row,
        column=1,
        value="INDICADORES PRINCIPALES",
    )
    ws.cell(row=next_row, column=1).fill = PatternFill(
        "solid",
        fgColor=RED,
    )
    ws.cell(row=next_row, column=1).font = Font(
        color=WHITE,
        bold=True,
    )

    metrics = [
        ("Colegios en cartera", summary["schools"]),
        ("Actividades del periodo", summary["activities"]),
        (
            "Oportunidades abiertas",
            summary["open_opportunities"],
        ),
        (
            "Oportunidades ganadas",
            summary["won_opportunities"],
        ),
        (
            "Oportunidades no concretadas",
            summary["lost_opportunities"],
        ),
        (
            "Unidades proyectadas",
            summary["projected_units"],
        ),
        (
            "Unidades adoptadas",
            summary["adopted_units"],
        ),
        (
            "Conversión de cierres (%)",
            summary["conversion_rate"],
        ),
        (
            "Colegios sin asesor",
            unassigned["schools"],
        ),
        (
            "Oportunidades sin asesor",
            unassigned["opportunities"],
        ),
    ]

    for offset, (label, value) in enumerate(metrics, start=1):
        row = next_row + offset
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=1).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 22


def _add_advisors_sheet(wb, advisor_report):
    ws = wb.create_sheet("02_Asesores")
    _style_title(
        ws,
        "REPORTE POR ASESOR",
        "Cartera, productividad y resultados comerciales",
    )

    headers = [
        "Asesor",
        "Equipo(s)",
        "Colegios",
        "Actividades",
        "Llamadas",
        "Visitas coordinadas",
        "Visitas en frío",
        "Presentaciones",
        "Reuniones",
        "Seguimientos",
        "Abiertas",
        "Ganadas",
        "No concretadas",
        "Unidades proyectadas",
        "Unidades adoptadas",
        "Conversión %",
    ]
    rows = []

    for row in advisor_report["advisors"]:
        counts = row["activity_counts"]
        rows.append(
            [
                row["advisor_name"],
                ", ".join(row["teams"]),
                row["schools"],
                row["activities"],
                counts.get("call", 0),
                counts.get("visit", 0),
                counts.get("cold_visit", 0),
                counts.get("presentation", 0),
                counts.get("meeting", 0),
                counts.get("follow_up", 0),
                row["open_opportunities"],
                row["won_opportunities"],
                row["lost_opportunities"],
                row["projected_units"],
                row["adopted_units"],
                row["conversion_rate"],
            ]
        )

    _write_table(
        ws,
        start_row=4,
        headers=headers,
        rows=rows,
        widths=[
            24, 28, 11, 12, 10, 15, 13, 14,
            12, 13, 10, 10, 14, 17, 17, 12,
        ],
        percent_columns={16},
    )


def _add_editorials_sheet(wb, editorial_report):
    ws = wb.create_sheet("03_Editoriales")
    _style_title(
        ws,
        "REPORTE POR EDITORIAL",
        "Proyección y adopción por editorial",
    )

    headers = [
        "Editorial",
        "Colegios proyectados",
        "Productos proyectados",
        "Unidades proyectadas",
        "Valor referencial proyección",
        "Colegios adoptados",
        "Productos adoptados",
        "Unidades adoptadas",
        "Valor adoptado P.IE",
        "Conversión unidades %",
    ]
    rows = [
        [
            row["editorial"],
            row["projected_schools"],
            row["projected_products"],
            row["projected_units"],
            float(row["projected_reference_value"]),
            row["adopted_schools"],
            row["adopted_products"],
            row["adopted_units"],
            float(row["adopted_value"]),
            row["unit_conversion_rate"],
        ]
        for row in editorial_report["editorials"]
    ]

    _write_table(
        ws,
        start_row=4,
        headers=headers,
        rows=rows,
        widths=[
            28, 18, 18, 18, 24, 18, 18, 17, 21, 19,
        ],
        currency_columns={5, 9},
        percent_columns={10},
    )


def _add_schools_sheet(wb, school_report):
    ws = wb.create_sheet("04_Colegios")
    _style_title(
        ws,
        "REPORTE POR COLEGIO",
        "Estado comercial de la cartera institucional",
    )

    headers = [
        "Colegio",
        "Departamento",
        "Provincia",
        "Distrito",
        "Equipo",
        "Asesor",
        "Abiertas",
        "Ganadas",
        "No concretadas",
        "Unidades proyectadas",
        "Unidades adoptadas",
        "Última actividad",
    ]
    rows = [
        [
            row["school_name"],
            row["department"],
            row["province"],
            row["district"],
            row["team"],
            row["advisor"],
            row["open_opportunities"],
            row["won_opportunities"],
            row["lost_opportunities"],
            row["projected_units"],
            row["adopted_units"],
            _safe_local_datetime(row["last_activity_at"]),
        ]
        for row in school_report["schools"]
    ]

    _write_table(
        ws,
        start_row=4,
        headers=headers,
        rows=rows,
        widths=[
            34, 15, 17, 17, 24, 24, 10, 10,
            14, 18, 17, 20,
        ],
        date_columns={12},
    )


def _add_opportunities_sheet(wb, opportunity_report):
    ws = wb.create_sheet("05_Oportunidades")
    _style_title(
        ws,
        "REPORTE DE OPORTUNIDADES",
        "Pipeline, proyección, cotización y adopción",
    )

    headers = [
        "Colegio",
        "Campaña",
        "Equipo",
        "Asesor",
        "Etapa",
        "Línea comercial",
        "Unidades proyectadas",
        "Cotización",
        "Versión",
        "Adopción",
        "Unidades adoptadas",
        "Última actividad",
        "Fecha de cierre",
    ]
    rows = [
        [
            row["school_name"],
            row["campaign"],
            row["team"],
            row["advisor"],
            row["stage"],
            row["commercial_line"],
            row["projected_units"],
            row["quotation_status"],
            row["quotation_version"],
            row["adoption_status"],
            row["adopted_units"],
            _safe_local_datetime(row["last_activity_at"]),
            _safe_local_datetime(row["closed_at"]),
        ]
        for row in opportunity_report["opportunities"]
    ]

    _write_table(
        ws,
        start_row=4,
        headers=headers,
        rows=rows,
        widths=[
            34, 24, 24, 24, 24, 18, 18,
            16, 10, 14, 17, 20, 20,
        ],
        date_columns={12, 13},
    )


def _add_activities_sheet(wb, activity_report):
    ws = wb.create_sheet("06_Actividades")
    _style_title(
        ws,
        "REPORTE DE ACTIVIDADES",
        "Evidencia de gestión comercial registrada",
    )

    headers = [
        "Fecha",
        "Asesor",
        "Colegio",
        "Contacto",
        "Tipo",
        "Resumen",
        "Resultado",
        "Oportunidad",
        "Ubicación registrada",
        "Importante",
    ]
    rows = [
        [
            _safe_local_datetime(row["occurred_at"]),
            row["advisor"],
            row["school"],
            row["contact"],
            row["activity_type"],
            row["summary"],
            row["result"],
            row["opportunity"],
            "Sí" if row["has_location"] else "No",
            "Sí" if row["is_important"] else "No",
        ]
        for row in activity_report["activities"]
    ]

    _write_table(
        ws,
        start_row=4,
        headers=headers,
        rows=rows,
        widths=[
            20, 24, 32, 24, 20, 34, 44, 32, 19, 12,
        ],
        date_columns={1},
    )


def _add_profitability_sheet(wb, editorial_report):
    ws = wb.create_sheet("07_Rentabilidad_Interna")
    _style_title(
        ws,
        "ANÁLISIS COMERCIAL INTERNO",
        (
            "Contribución estimada sobre adopciones confirmadas. "
            "No representa utilidad neta contable."
        ),
    )

    ws["A3"] = (
        "Información sensible: uso interno de Administración y "
        "Jefatura Comercial."
    )
    ws["A3"].fill = PatternFill("solid", fgColor=AMBER)
    ws["A3"].font = Font(bold=True, color=BLACK)

    headers = [
        "Editorial",
        "Unidades adoptadas",
        "Valor adoptado P.IE",
        "Costo editorial total",
        "Incentivos / comisiones",
        "Contribución estimada",
        "Margen %",
        "Contribución por unidad",
    ]
    rows = [
        [
            row["editorial"],
            row["adopted_units"],
            float(row["adopted_value"]),
            float(row["supplier_cost_total"]),
            float(row["incentive_total"]),
            float(row["contribution_total"]),
            row["margin_percent"],
            float(row["contribution_per_unit"]),
        ]
        for row in editorial_report["editorials"]
    ]

    _write_table(
        ws,
        start_row=5,
        headers=headers,
        rows=rows,
        widths=[28, 18, 20, 21, 22, 22, 12, 23],
        currency_columns={3, 4, 5, 6, 8},
        percent_columns={7},
    )


def build_crm_report_workbook(
    *,
    user,
    date_from,
    date_to,
    campaign_id=None,
    team_id=None,
    owner_id=None,
    include_profitability=False,
):
    advisor_report = build_advisor_commercial_report(
        user=user,
        date_from=date_from,
        date_to=date_to,
        campaign_id=campaign_id,
        team_id=team_id,
        owner_id=owner_id,
    )
    editorial_report = build_editorial_commercial_report(
        user=user,
        campaign_id=campaign_id,
        team_id=team_id,
        owner_id=owner_id,
    )
    school_report = build_school_commercial_report(
        user=user,
        campaign_id=campaign_id,
        team_id=team_id,
        owner_id=owner_id,
    )
    opportunity_report = build_opportunity_commercial_report(
        user=user,
        campaign_id=campaign_id,
        team_id=team_id,
        owner_id=owner_id,
    )
    activity_report = build_activity_commercial_report(
        user=user,
        date_from=date_from,
        date_to=date_to,
        campaign_id=campaign_id,
        team_id=team_id,
        owner_id=owner_id,
    )

    labels = _filter_labels(
        user=user,
        campaign_id=campaign_id,
        team_id=team_id,
        owner_id=owner_id,
    )

    wb = Workbook()
    wb.properties.title = "Reporte Comercial CRM - Book Express"
    wb.properties.subject = "Reportería comercial consolidada"
    wb.properties.creator = "Book Express"

    _add_summary_sheet(
        wb,
        advisor_report=advisor_report,
        labels=labels,
        date_from=date_from,
        date_to=date_to,
    )
    _add_advisors_sheet(wb, advisor_report)
    _add_editorials_sheet(wb, editorial_report)
    _add_schools_sheet(wb, school_report)
    _add_opportunities_sheet(wb, opportunity_report)
    _add_activities_sheet(wb, activity_report)

    if include_profitability:
        _add_profitability_sheet(wb, editorial_report)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    return stream
