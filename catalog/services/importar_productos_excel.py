from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils.text import slugify
from openpyxl import load_workbook

from catalog.models import (
    Provider,
    Level,
    Grade,
    Area,
    Series,
    ProductType,
    Product,
    ProductPrice,
    CargaExcel,
    CargaExcelDetalle,
)


COLUMNAS_OBLIGATORIAS = [
    "proveedor_editorial",
    "nombre",
    "tipo_producto",
    "anio_catalogo",
    "consultar_precio",
    "activo",
]


def limpiar_texto(valor):
    if valor is None:
        return ""
    return str(valor).strip()
def obtener_o_crear_por_slug(modelo, nombre):
    """
    Busca un registro por slug antes de crearlo.
    Evita errores cuando existen nombres parecidos como:
    '3 años', '3 Años', '3 AÑOS', etc.
    """

    nombre = limpiar_texto(nombre)

    if not nombre:
        return None, False

    slug = slugify(nombre)

    objeto = modelo.objects.filter(slug=slug).first()

    if objeto:
        return objeto, False

    return modelo.objects.get_or_create(
        name=nombre,
        defaults={
            "slug": slug,
        }
    )


def obtener_o_crear_serie(provider, nombre):
    """
    Busca o crea una serie asociada a un proveedor.
    """

    nombre = limpiar_texto(nombre)

    if not nombre:
        return None, False

    slug = slugify(f"{provider.name}-{nombre}")

    objeto = Series.objects.filter(
        provider=provider,
        slug=slug
    ).first()

    if objeto:
        return objeto, False

    return Series.objects.get_or_create(
        provider=provider,
        name=nombre,
        defaults={
            "slug": slug,
        }
    )

def normalizar_columna(valor):
    return limpiar_texto(valor).lower()


def convertir_booleano(valor):
    valor = limpiar_texto(valor).upper()

    if valor in ["SI", "SÍ", "TRUE", "1", "ACTIVO"]:
        return True

    if valor in ["NO", "FALSE", "0", "INACTIVO"]:
        return False

    return None


def convertir_decimal(valor):
    if valor is None or valor == "":
        return None

    try:
        return Decimal(str(valor).replace(",", ".")).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def obtener_valor(fila, columnas, nombre_columna):
    indice = columnas.get(nombre_columna)

    if indice is None:
        return None

    return fila[indice]


def mapear_disponibilidad(valor):
    valor = limpiar_texto(valor).upper()

    if valor in ["DISPONIBLE", "AVAILABLE", "SI", "SÍ"]:
        return "available"

    if valor in ["STOCK LIMITADO", "LIMITADO", "LIMITED"]:
        return "limited"

    if valor in ["AGOTADO", "OUT OF STOCK"]:
        return "out_of_stock"

    return "on_request"


def validar_fila(data):
    errores = []

    for columna in COLUMNAS_OBLIGATORIAS:
        if not limpiar_texto(data.get(columna)):
            errores.append(f"La columna '{columna}' es obligatoria.")

    if not limpiar_texto(data.get("codigo_producto")) and not limpiar_texto(data.get("sku")):
        errores.append("Debe existir codigo_producto o sku.")

    try:
        int(data.get("anio_catalogo"))
    except Exception:
        errores.append("El anio_catalogo debe ser numérico.")

    activo = convertir_booleano(data.get("activo"))
    consultar_precio = convertir_booleano(data.get("consultar_precio"))

    if activo is None:
        errores.append("La columna activo debe tener SI o NO.")

    if consultar_precio is None:
        errores.append("La columna consultar_precio debe tener SI o NO.")

    precio_referencial = convertir_decimal(data.get("precio_referencial"))

    if consultar_precio is False and precio_referencial is None:
        errores.append("Si consultar_precio es NO, debe existir precio_referencial.")

    return errores


def construir_clave_producto(data):
    proveedor = limpiar_texto(data.get("proveedor_editorial")).upper()
    codigo = limpiar_texto(data.get("codigo_producto")).upper()
    sku = limpiar_texto(data.get("sku")).upper()
    nombre = limpiar_texto(data.get("nombre")).upper()

    if codigo:
        return f"{proveedor}|CODIGO|{codigo}"

    if sku:
        return f"{proveedor}|SKU|{sku}"

    return f"{proveedor}|NOMBRE|{nombre}"


def buscar_producto_existente(data):
    proveedor_nombre = limpiar_texto(data.get("proveedor_editorial"))
    codigo = limpiar_texto(data.get("codigo_producto"))
    sku = limpiar_texto(data.get("sku"))
    nombre = limpiar_texto(data.get("nombre"))

    producto = None

    if codigo:
        producto = Product.objects.filter(
            provider__name__iexact=proveedor_nombre,
            code__iexact=codigo,
        ).first()

    if not producto and sku:
        producto = Product.objects.filter(
            provider__name__iexact=proveedor_nombre,
            sku__iexact=sku,
        ).first()

    if not producto and nombre:
        producto = Product.objects.filter(
            provider__name__iexact=proveedor_nombre,
            name__iexact=nombre,
        ).first()

    return producto


def leer_datos_excel(ruta_archivo):
    workbook = load_workbook(ruta_archivo, data_only=True)

    if "Catalogo_Maestro_2026" in workbook.sheetnames:
        sheet = workbook["Catalogo_Maestro_2026"]
    else:
        sheet = workbook.active

    columnas = {}

    for indice, celda in enumerate(sheet[1]):
        if celda.value:
            columnas[normalizar_columna(celda.value)] = indice

    filas = []

    for numero_fila, fila in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if all(valor is None for valor in fila):
            continue

        data = {
            "proveedor_editorial": limpiar_texto(obtener_valor(fila, columnas, "proveedor_editorial")),
            "codigo_producto": limpiar_texto(obtener_valor(fila, columnas, "codigo_producto")),
            "sku": limpiar_texto(obtener_valor(fila, columnas, "sku")),
            "nombre": limpiar_texto(obtener_valor(fila, columnas, "nombre")),
            "nombre_original": limpiar_texto(obtener_valor(fila, columnas, "nombre_original")),
            "unidad": limpiar_texto(obtener_valor(fila, columnas, "unidad")),
            "categoria_original": limpiar_texto(obtener_valor(fila, columnas, "categoria_original")),
            "subcategoria_original": limpiar_texto(obtener_valor(fila, columnas, "subcategoria_original")),
            "nivel": limpiar_texto(obtener_valor(fila, columnas, "nivel")),
            "grado": limpiar_texto(obtener_valor(fila, columnas, "grado")),
            "area": limpiar_texto(obtener_valor(fila, columnas, "area")),
            "serie": limpiar_texto(obtener_valor(fila, columnas, "serie")),
            "tipo_producto": limpiar_texto(obtener_valor(fila, columnas, "tipo_producto")),
            "anio_catalogo": obtener_valor(fila, columnas, "anio_catalogo"),
            "precio_costo": obtener_valor(fila, columnas, "precio_costo"),
            "precio_referencial": obtener_valor(fila, columnas, "precio_referencial"),
            "consultar_precio": limpiar_texto(obtener_valor(fila, columnas, "consultar_precio")),
            "disponibilidad": limpiar_texto(obtener_valor(fila, columnas, "disponibilidad")),
            "activo": limpiar_texto(obtener_valor(fila, columnas, "activo")),
            "descripcion": limpiar_texto(obtener_valor(fila, columnas, "descripcion")),
            "observaciones": limpiar_texto(obtener_valor(fila, columnas, "observaciones")),
        }

        filas.append((numero_fila, data))

    return filas


def crear_vista_previa_productos(archivo, anio_catalogo):
    carga = CargaExcel.objects.create(
        archivo=archivo,
        anio_catalogo=anio_catalogo,
        estado="PENDIENTE",
    )

    filas = leer_datos_excel(carga.archivo.path)

    claves_excel = set()

    total_filas = 0
    total_nuevos = 0
    total_actualizados = 0
    total_errores = 0

    for numero_fila, data in filas:
        total_filas += 1

        if not data.get("anio_catalogo"):
            data["anio_catalogo"] = anio_catalogo

        errores = validar_fila(data)

        clave = construir_clave_producto(data)

        if clave in claves_excel:
            errores.append("Producto duplicado dentro del mismo Excel.")

        claves_excel.add(clave)

        if errores:
            accion = "ERROR"
            total_errores += 1
        else:
            producto_existente = buscar_producto_existente(data)

            if producto_existente:
                accion = "ACTUALIZAR"
                total_actualizados += 1
            else:
                accion = "NUEVO"
                total_nuevos += 1

        CargaExcelDetalle.objects.create(
            carga=carga,
            numero_fila=numero_fila,
            codigo_producto=data.get("codigo_producto"),
            sku=data.get("sku"),
            nombre=data.get("nombre"),
            proveedor_editorial=data.get("proveedor_editorial"),
            accion=accion,
            errores=errores,
            datos_originales=data,
            procesado=False,
        )

    carga.total_filas = total_filas
    carga.total_nuevos = total_nuevos
    carga.total_actualizados = total_actualizados
    carga.total_errores = total_errores
    carga.estado = "VALIDADO" if total_errores == 0 else "ERROR"
    carga.save()

    return carga


@transaction.atomic
def confirmar_importacion_productos(carga_id):
    carga = CargaExcel.objects.get(id=carga_id)

    if carga.total_errores > 0:
        raise ValueError("No se puede importar. Existen filas con errores.")

    detalles = carga.detalles.filter(procesado=False)

    for detalle in detalles:
        data = detalle.datos_originales

        provider, _ = obtener_o_crear_por_slug(
        Provider,
        data["proveedor_editorial"]
        )

        product_type, _ = obtener_o_crear_por_slug(
        ProductType,
        data["tipo_producto"]
        )

        level = None
        if data.get("nivel"):
            level, _ = obtener_o_crear_por_slug(
                Level,
                data["nivel"]
            )

        grade = None
        if data.get("grado"):
            grade, _ = obtener_o_crear_por_slug(
                Grade,
                data["grado"]
            )

        area = None
        if data.get("area"):
            area, _ = obtener_o_crear_por_slug(
                Area,
                data["area"]
            )

        series = None
        if data.get("serie"):
            series, _ = obtener_o_crear_serie(
                provider,
                data["serie"]
            )

        is_active = convertir_booleano(data.get("activo"))
        consult_price = convertir_booleano(data.get("consultar_precio"))

        product = buscar_producto_existente(data)

        if product is None:
            product = Product.objects.create(
                provider=provider,
                code=data.get("codigo_producto"),
                sku=data.get("sku"),
                name=data.get("nombre"),
                series=series,
                level=level,
                grade=grade,
                area=area,
                product_type=product_type,
                description=data.get("descripcion"),
                is_active=is_active,
            )
        else:
            product.provider = provider
            product.code = data.get("codigo_producto")
            product.sku = data.get("sku")
            product.name = data.get("nombre")
            product.series = series
            product.level = level
            product.grade = grade
            product.area = area
            product.product_type = product_type
            product.description = data.get("descripcion")
            product.is_active = is_active
            product.save()

        precio_referencial = convertir_decimal(data.get("precio_referencial"))
        precio_costo = convertir_decimal(data.get("precio_costo"))

        ProductPrice.objects.update_or_create(
            product=product,
            year=int(data["anio_catalogo"]),
            campaign="Campaña escolar",
            defaults={
                "cost_price": precio_costo,
                "price": precio_referencial,
                "show_price": not consult_price and precio_referencial is not None,
                "consult_price": consult_price,
                "availability": mapear_disponibilidad(data.get("disponibilidad")),
                "is_active": is_active,
            },
        )

        detalle.procesado = True
        detalle.save()

    carga.estado = "IMPORTADO"
    carga.save()

    return carga